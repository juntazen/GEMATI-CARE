from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips


def _font(run, size: float = 9, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Times New Roman"
    )
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _replace_paragraph(paragraph, text: str, size: float, **font_kwargs) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    _font(run, size, **font_kwargs)


def _remove_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def _remove_body_siblings_between(start, end) -> None:
    """Remove every body-level node between two template sentinels.

    python-docx omits structured-document tags (w:sdt) from ``paragraphs``.
    Traversing XML siblings prevents hidden sample bibliography blocks from
    surviving in the rendered manuscript while preserving both section breaks.
    """
    node = start.getnext()
    while node is not None and node is not end:
        following = node.getnext()
        node.getparent().remove(node)
        node = following


def _clear_editor_placeholders(document: Document) -> None:
    targets = ("[Caption completed by the editor]", "How to Cite:", "Permalink/DOI:")
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        cleaned = paragraph.text
        for target in targets:
            cleaned = cleaned.replace(target, "")
        if cleaned != paragraph.text:
            _replace_paragraph(paragraph, cleaned.strip(), 8)


def _extract(markdown: str) -> dict:
    lines = markdown.splitlines()
    title = next(line[2:].strip() for line in lines if line.startswith("# "))
    abstract_heading = "Abstrak" if "## Abstrak" in lines else "Abstract"
    abstract_index = lines.index(f"## {abstract_heading}")
    keyword_prefix = "Kata kunci:" if abstract_heading == "Abstrak" else "Keywords:"
    keyword_index = next(i for i, line in enumerate(lines) if line.startswith(keyword_prefix))
    front = [line.strip() for line in lines[1:abstract_index] if line.strip()]
    abstract = " ".join(line.strip() for line in lines[abstract_index + 1 : keyword_index] if line.strip())
    body_index = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    return {
        "language": "id" if abstract_heading == "Abstrak" else "en",
        "title": title,
        "authors": front[0] if front else "[Authors]",
        "affiliation": front[1] if len(front) > 1 else "[Affiliations]",
        "email": front[2] if len(front) > 2 else "[Corresponding author email]",
        "abstract_heading": abstract_heading,
        "abstract": abstract,
        "keywords": lines[keyword_index],
        "body_lines": lines[body_index:],
    }


def _move_before(element, sentinel) -> None:
    sentinel.addprevious(element)


def _new_paragraph(document: Document, sentinel, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    _move_before(paragraph._p, sentinel)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    _font(run, 9)
    return paragraph


def _set_horizontal_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "insideH", "bottom", "left", "right", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single" if edge in {"top", "insideH", "bottom"} else "nil")
        if edge in {"top", "insideH", "bottom"}:
            node.set(qn("w:sz"), "6")
            node.set(qn("w:color"), "000000")


def _compact_cell(cell, width, font_size: float) -> None:
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(int(width.twips)))
    tc_width.set(qn("w:type"), "dxa")
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)


def _insert_table(
    document: Document,
    sentinel,
    lines: list[str],
    number: int,
    language: str,
    caption_text: str | None = None,
) -> None:
    rows = [[item.strip() for item in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    caption = _new_paragraph(
        document,
        sentinel,
        ("Tabel" if language == "id" else "Table") + f" {number}. "
        + (caption_text or ("Titik operasi keselamatan–utilitas" if language == "id"
                            else "Safety–utility operating points")),
        "Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        _font(run, 8)
    table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    _move_before(table._tbl, sentinel)
    table.autofit = False
    column_count = max(len(row) for row in rows)
    column_width = Inches(2.92 / column_count)
    font_size = 6.2 if column_count >= 6 else 6.8 if column_count >= 5 else 7.2
    _set_horizontal_borders(table)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            _font(run, font_size, bold=row_index == 0)
            _compact_cell(cell, column_width, font_size)
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))


def _insert_body(
    document: Document,
    sentinel,
    body_lines: list[str],
    language: str,
    architecture_path: Path | None,
) -> None:
    table_number = 0
    pending_table_caption: str | None = None
    figure_number = 0
    figure_map = {
        "architecture": architecture_path,
        "reliability": architecture_path.parent.parent.parent / "results" / "manuscript_figures" / "reliability_tfidf_final.png" if architecture_path else None,
        "risk_coverage": architecture_path.parent.parent.parent / "results" / "manuscript_figures" / "conformal_tradeoff_tfidf_final.png" if architecture_path else None,
    }
    index = 0
    while index < len(body_lines):
        line = body_lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("[TABLE:") and line.endswith("]"):
            pending_table_caption = line[7:-1].strip()
            index += 1
            continue
        if line.startswith("[FIGURE:") and line.endswith("]"):
            payload = line[8:-1]
            key, _, caption_text = payload.partition("|")
            image_path = figure_map.get(key.strip())
            if image_path is None or not image_path.exists():
                raise FileNotFoundError(f"Figure marker refers to a missing image: {key}")
            figure_number += 1
            image_p = document.add_paragraph()
            _move_before(image_p._p, sentinel)
            image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_p.add_run().add_picture(str(image_path), width=Inches(3.0))
            cap = _new_paragraph(
                document,
                sentinel,
                ("Gambar" if language == "id" else "Figure")
                + f" {figure_number}. " + caption_text.strip(),
                "Caption",
            )
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs:
                _font(run, 8)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(body_lines) and body_lines[index].strip().startswith("|"):
                table_lines.append(body_lines[index].strip())
                index += 1
            table_number += 1
            _insert_table(document, sentinel, table_lines, table_number, language, pending_table_caption)
            pending_table_caption = None
            continue
        if line.startswith("## "):
            paragraph = _new_paragraph(document, sentinel, line[3:].strip(), "Heading 1")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                _font(run, 10, bold=True)
        elif line.startswith("### "):
            paragraph = _new_paragraph(document, sentinel, line[4:].strip(), "Heading 2")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                _font(run, 9, bold=True)
        elif re.match(r"^\[\d+\]", line):
            style = "Rujukan" if "Rujukan" in [s.name for s in document.styles] else "Normal"
            paragraph = _new_paragraph(document, sentinel, line, style)
            paragraph.paragraph_format.left_indent = Inches(0.16)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            for run in paragraph.runs:
                _font(run, 8)
        else:
            _new_paragraph(document, sentinel, line)
        index += 1


def _validate_docx(path: Path) -> dict:
    document = Document(path)
    section_columns = []
    for section in document.sections:
        cols = section._sectPr.find(qn("w:cols"))
        section_columns.append(int(cols.get(qn("w:num"), "1")) if cols is not None else 1)
    full_text = "\n".join(document._element.xpath(".//w:t/text()"))
    title = document.core_properties.title or ""
    abstract_text = ""
    for i, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip() in {"Abstrak", "Abstract"} and i + 1 < len(document.paragraphs):
            abstract_text = document.paragraphs[i + 1].text
            break
    checks = {
        "template_sections_valid": len(document.sections) >= 2,
        "body_section_two_columns": len(section_columns) >= 2 and section_columns[1] == 2,
        "title_at_most_12_words": len(title.split()) <= 12,
        "abstract_at_most_250_words": len(abstract_text.split()) <= 250,
        "at_least_25_references": len(re.findall(r"^\[\d+\]", full_text, flags=re.MULTILINE)) >= 25,
        "template_instruction_removed": "The introduction section sets the stage" not in full_text,
        "sample_bibliography_removed": "Writing a Scientific Review Article" not in full_text,
        "no_stale_header": all("Author1" not in p.text and "(2024)" not in p.text for s in document.sections for p in s.header.paragraphs),
        "editor_placeholders_removed": all(x not in full_text for x in ("How to Cite:", "Permalink/DOI:", "Caption completed by the editor")),
        "page_field_preserved": any("PAGE" in footer._element.xml for section in document.sections for footer in [section.footer]),
    }
    return {
        "path": str(path),
        "sections": len(document.sections),
        "section_columns": section_columns,
        "title_words": len(title.split()),
        "abstract_words": len(abstract_text.split()),
        "checks": checks,
        "pass": all(checks.values()),
    }


def create_docx_from_template(
    template_path: Path,
    markdown_path: Path,
    output_path: Path,
    architecture_path: Path | None = None,
) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"Template tidak ditemukan: {template_path}")
    content = _extract(markdown_path.read_text(encoding="utf-8"))
    document = Document(template_path)
    _clear_editor_placeholders(document)
    short_header = "J. Zeniarja et al. — Jurnal RESTI"
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                _replace_paragraph(paragraph, short_header, 8)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    original_paragraphs = list(document.paragraphs)
    if len(original_paragraphs) < 71:
        raise RuntimeError("Struktur TemplateRESTI2026 tidak dikenali.")

    _replace_paragraph(original_paragraphs[1], content["title"], 15)
    original_paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _replace_paragraph(original_paragraphs[2], content["authors"], 10)
    _replace_paragraph(original_paragraphs[3], content["affiliation"], 9)
    _replace_paragraph(original_paragraphs[4], "", 9)
    _replace_paragraph(original_paragraphs[5], content["email"], 9)
    _replace_paragraph(original_paragraphs[7], content["abstract_heading"], 10, bold=True)
    _replace_paragraph(original_paragraphs[8], content["abstract"], 9, italic=True)
    original_paragraphs[8].style = document.styles["Normal"]
    original_paragraphs[8].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _replace_paragraph(original_paragraphs[9], content["keywords"], 9)

    sentinel = original_paragraphs[70]._p
    _remove_body_siblings_between(original_paragraphs[16]._p, sentinel)
    _insert_body(document, sentinel, content["body_lines"], content["language"], architecture_path)
    _remove_paragraph(original_paragraphs[70])
    body_cols = document.sections[-1]._sectPr.find(qn("w:cols"))
    if body_cols is None:
        body_cols = OxmlElement("w:cols")
        document.sections[-1]._sectPr.append(body_cols)
    body_cols.set(qn("w:num"), "2")
    body_cols.set(qn("w:space"), "227")

    document.core_properties.title = content["title"]
    document.core_properties.author = "Junta Zeniarja; Erwin Yudi Hidayat; Egia Rosi Subhiyakto"
    document.core_properties.subject = "GEMATI-CARE RESTI research manuscript draft"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def generate_bilingual_docx(root: Path, template_path: Path) -> list[Path]:
    progress = json.loads((root / "results" / "extended_progress.json").read_text(encoding="utf-8"))
    if progress.get("status") != "complete":
        raise RuntimeError("Suite extended belum complete; DOCX tidak dibuat.")
    architecture = root / "docs" / "figures" / "arsitektur_gemati_care.png"
    outputs = [
        create_docx_from_template(
            template_path,
            root / "paper" / "DRAFT_PAPER_RESTI_ID.md",
            root / "paper" / "DRAFT_PAPER_RESTI_ID.docx",
            architecture,
        ),
        create_docx_from_template(
            template_path,
            root / "paper" / "DRAFT_PAPER_RESTI_EN.md",
            root / "paper" / "DRAFT_PAPER_RESTI_EN.docx",
            architecture,
        ),
    ]
    audits = [_validate_docx(path) for path in outputs]
    (root / "results" / "DOCX_AUDIT.json").write_text(
        json.dumps(audits, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    if not all(item["pass"] for item in audits):
        raise RuntimeError("DOCX dibuat tetapi gagal structural audit; lihat results/DOCX_AUDIT.json")
    return outputs
